from datetime import datetime, timedelta
import secrets
import csv
import io

from pydantic import BaseModel
from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import AuthContext, get_current_auth, get_session_dep
from app.api.deps import has_permission
from app.db.models import AuditLog, Contact, Conversation, Customer, Inbox, Lead, LeadActivity, LeadArea, LeadProfession, LeadSector, LeadSource, LeadStage, LeadShareLink, Message, Opportunity, Organization, SalesOrder, Task, UserLLMConfig
from app.inbox.security import decrypt_channel_config
from app.schemas.common import PaginationMeta
from app.schemas.customer import CustomerOut
from app.schemas.lead import (
    LeadActivityCreate,
    LeadActivityOut,
    LeadActivityUpdate,
    LeadAnalyticsSummary,
    LeadConvertPayload,
    LeadCreate,
    LeadListResponse,
    LeadOut,
    LeadShareCreate,
    LeadShareOut,
    LeadUpdate,
    LeadTimelineItem,
)
from app.core.config import settings
from app.services.ai_convert import convert_notes_to_lead
from app.services.intelligence import evaluate_intelligence_alerts, record_preference_history, record_stage_history
from app.services.lead_capture import upsert_lead_from_inbound_message


router = APIRouter()

async def _get_lead_or_404(lead_id: str, session: AsyncSession, auth: AuthContext = None) -> Lead:
    lead = await session.get(Lead, lead_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    if auth and lead.organization_id and lead.organization_id != auth.org_id:
        raise HTTPException(status_code=404, detail="Lead not found")
    if auth:
        from app.api.deps import apply_tenant_filters
        stmt = apply_tenant_filters(select(Lead).where(Lead.id == lead_id), auth, Lead)
        res = await session.execute(stmt)
        if not res.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Lead not found")
    return lead


async def _populate_lead_assignments(lead_obj: Lead, session: AsyncSession) -> list[str]:
    from app.db.models import Assignment
    res = await session.execute(select(Assignment.user_id).where(Assignment.lead_id == lead_obj.id))
    return list(res.scalars().all())


async def _validate_config_ids(changes: dict, session: AsyncSession) -> None:
    checks = {
        "lead_source_id": LeadSource,
        "lead_stage_id": LeadStage,
        "lead_sector_id": LeadSector,
        "lead_area_id": LeadArea,
        "lead_profession_id": LeadProfession,
    }
    for field_name, model in checks.items():
        value = changes.get(field_name)
        if value is None or value == "":
            continue
        entity = await session.get(model, value)
        if not entity or not entity.is_active:
            raise HTTPException(status_code=400, detail=f"Invalid {field_name}")


def _normalize_tags(tags: list[str] | None) -> list[str] | None:
    if tags is None:
        return None
    normalized = [tag.strip().lower() for tag in tags if isinstance(tag, str) and tag.strip()]
    unique: list[str] = []
    seen = set()
    for tag in normalized:
        if tag not in seen:
            seen.add(tag)
            unique.append(tag)
    return unique


def _normalize_emails(emails: list[str] | None) -> list[str]:
    if not emails:
        return []
    normalized = [email.strip().lower() for email in emails if isinstance(email, str) and email.strip()]
    unique: list[str] = []
    seen = set()
    for email in normalized:
        if email not in seen:
            seen.add(email)
            unique.append(email)
    return unique


async def _can_view_all_leads(auth: AuthContext, session: AsyncSession) -> bool:
    return auth.role == "super_admin" or await has_permission(session, auth.user_id, auth, "leads.manage")


async def _can_view_leads(auth: AuthContext, session: AsyncSession) -> bool:
    return auth.role == "super_admin" or await has_permission(session, auth.user_id, auth, "leads.view") or await _can_view_all_leads(auth, session)


def _is_assigned_lead(lead: Lead, auth: AuthContext) -> bool:
    return lead.assigned_user_id == auth.user_id or lead.assigned_agent_id == auth.user_id


@router.get("", response_model=LeadListResponse)
async def list_leads(
    status_filter: str | None = Query(default=None, alias="status"),
    stage_id: str | None = Query(default=None),
    source_id: str | None = Query(default=None),
    tag: str | None = Query(default=None),
    priority: str | None = Query(default=None),
    search: str | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    assigned_user_id: str | None = Query(default=None),
    branch_id: str | None = Query(default=None),
    quick_filter: str | None = Query(default=None),
    sort: str = Query(default="desc", pattern="^(asc|desc)$"),
    limit: int = Query(default=20, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadListResponse:
    assigned_scope_requested = quick_filter == "assigned_to_me"
    assigned_user_lookup_requested = bool(assigned_user_id)
    if assigned_user_lookup_requested and auth.role != "super_admin" and not await _can_view_all_leads(auth, session):
        raise HTTPException(status_code=403, detail="Permission denied")
    if not assigned_scope_requested and not await _can_view_leads(auth, session):
        raise HTTPException(status_code=403, detail="Permission denied")

    filters = []

    if status_filter and status_filter != "all":
        filters.append(Lead.status == status_filter)
    if branch_id:
        filters.append(Lead.branch_id == branch_id)
    if stage_id:
        filters.append(Lead.lead_stage_id == stage_id)
    if source_id:
        filters.append(Lead.lead_source_id == source_id)
    if tag:
        tag_value = tag.strip().lower()
        if tag_value:
            filters.append(Lead.tags.contains([tag_value]))
    if priority and priority != "all":
        filters.append(Lead.priority == priority.lower())
    if search:
        needle = f"%{search.strip()}%"
        filters.append(
            or_(
                Lead.company_name.ilike(needle),
                Lead.contact_person.ilike(needle),
                Lead.phone.ilike(needle),
                Lead.email.ilike(needle),
            )
        )
    if start_date:
        filters.append(Lead.created_at >= start_date)
    if end_date:
        filters.append(Lead.created_at <= end_date)
    if assigned_user_lookup_requested:
        filters.append(or_(Lead.assigned_user_id == assigned_user_id, Lead.assigned_agent_id == assigned_user_id))
    if assigned_scope_requested:
        filters.append(or_(Lead.assigned_user_id == auth.user_id, Lead.assigned_agent_id == auth.user_id))
    elif quick_filter == "untouched":
        filters.append(Lead.untouched.is_(True))
    elif quick_filter == "followups_due":
        filters.append(and_(Lead.follow_up_date.is_not(None), Lead.follow_up_date <= func.now()))

    query = select(Lead)
    count_query = select(func.count(Lead.id))

    if filters:
        query = query.where(*filters)
        count_query = count_query.where(*filters)

    from app.api.deps import apply_tenant_filters
    query = apply_tenant_filters(query, auth, Lead)
    count_query = apply_tenant_filters(count_query, auth, Lead)

    order_column = Lead.created_at.asc() if sort == "asc" else Lead.created_at.desc()
    query = query.order_by(order_column).limit(limit).offset(offset)
    rows = (await session.execute(query)).scalars().all()
    total = (await session.execute(count_query)).scalar_one()

    data_out = []
    for r in rows:
        l_out = LeadOut.model_validate(r)
        l_out.assigned_user_ids = await _populate_lead_assignments(r, session)
        data_out.append(l_out)

    return LeadListResponse(
        data=data_out,
        meta=PaginationMeta(total=total, limit=limit, offset=offset),
    )


@router.get("/analytics/summary", response_model=LeadAnalyticsSummary)
async def get_lead_analytics_summary(
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadAnalyticsSummary:
    from app.api.deps import apply_tenant_filters
    total = (await session.execute(apply_tenant_filters(select(func.count(Lead.id)), auth, Lead))).scalar_one()
    converted = (await session.execute(apply_tenant_filters(select(func.count(Lead.id)).where(Lead.converted_customer_id.is_not(None)), auth, Lead))).scalar_one()

    status_rows = (await session.execute(apply_tenant_filters(select(Lead.status, func.count(Lead.id)).group_by(Lead.status), auth, Lead))).all()
    source_rows = (await session.execute(apply_tenant_filters(select(Lead.source, func.count(Lead.id)).group_by(Lead.source), auth, Lead))).all()

    return LeadAnalyticsSummary(
        total=total,
        by_status={status_name or "unknown": count for status_name, count in status_rows},
        by_source={source_name or "unsourced": count for source_name, count in source_rows},
        converted=converted,
        conversion_rate=round((converted / total) * 100, 1) if total else 0,
    )


class AIConvertPayload(BaseModel):
    raw_notes: str


@router.post("", response_model=LeadOut, status_code=status.HTTP_201_CREATED)
async def create_lead(
    payload: LeadCreate,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadOut:
    data = payload.model_dump()
    data["organization_id"] = auth.org_id
    data["branch_id"] = auth.branch_id
    if data.get("priority"):
        data["priority"] = data["priority"].lower()
    if "tags" in data:
        data["tags"] = _normalize_tags(data.get("tags"))
    await _validate_config_ids(data, session)

    # Validate phone uniqueness for registration
    phone = data.get("phone")
    if phone and phone.strip():
        phone_stripped = phone.strip()
        q = select(Lead).where(Lead.phone == phone_stripped)
        res = await session.execute(q)
        existing = res.scalars().first()
        if existing:
            existing_name = existing.company_name or existing.contact_person or "Unnamed Lead"
            raise HTTPException(
                status_code=400,
                detail=f"Lead with this phone number already exists: {existing_name}",
            )

    assigned_user_ids = data.pop("assigned_user_ids", None)
    entity = Lead(**data)

    if entity.assigned_user_id or entity.assigned_agent_id or assigned_user_ids:
        from datetime import datetime, timezone, timedelta
        entity.assigned_at = datetime.now(timezone.utc)
        if entity.sla_duration_hours:
            entity.target_stage_by = entity.assigned_at + timedelta(hours=entity.sla_duration_hours)

    session.add(entity)
    await session.flush()

    if assigned_user_ids:
        from app.db.models import Assignment
        for uid in assigned_user_ids:
            session.add(Assignment(
                organization_id=auth.org_id,
                lead_id=entity.id,
                user_id=uid,
                assigned_by_id=auth.user_id
            ))

    await record_preference_history(
        session,
        workspace_id=auth.user_id,
        lead=entity,
        text=" ".join(filter(None, [entity.notes, entity.raw_note, entity.company_name, entity.address, entity.industry])),
        detected_from="lead_create",
    )
    await session.commit()
    await session.refresh(entity)
    l_out = LeadOut.model_validate(entity)
    l_out.assigned_user_ids = await _populate_lead_assignments(entity, session)
    return l_out


@router.post("/ai-convert", response_model=LeadCreate)
async def ai_convert_notes(
    payload: AIConvertPayload,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadCreate:
    api_key = None
    q = select(UserLLMConfig).where(
        UserLLMConfig.user_id == auth.user_id,
        UserLLMConfig.provider == "groq",
    )
    res = await session.execute(q)
    config_record = res.scalar_one_or_none()

    if config_record and config_record.encrypted_config:
        try:
            decrypted = decrypt_channel_config(config_record.encrypted_config)
            api_key = decrypted.get("api_key")
        except Exception:
            pass

    if not api_key:
        raise HTTPException(
            status_code=400,
            detail="Groq API Key not configured. Please add your key in Settings > AI & Automation.",
        )

    org_type_code = None
    if auth.org_id:
        from app.db.models import Organization, OrganizationType
        org = await session.get(Organization, auth.org_id)
        if org and org.organization_type_id:
            org_type = await session.get(OrganizationType, org.organization_type_id)
            if org_type:
                org_type_code = org_type.code

    result = await convert_notes_to_lead(payload.raw_notes, api_key, org_type_code)
    if not result:
        raise HTTPException(
            status_code=400,
            detail="AI extraction failed. Please verify your API key in Settings or fill in manually.",
        )
    return result


@router.get("/from-conversation/{conversation_id}", response_model=LeadOut)
async def get_lead_from_conversation(
    conversation_id: str,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadOut:
    from app.api.deps import apply_tenant_filters
    lead = (
        await session.execute(
            apply_tenant_filters(
                select(Lead)
                .where(Lead.conversation_id == conversation_id)
                .order_by(Lead.updated_at.desc()),
                auth,
                Lead
            ).limit(1)
        )
    ).scalar_one_or_none()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found for conversation")
    l_out = LeadOut.model_validate(lead)
    l_out.assigned_user_ids = await _populate_lead_assignments(lead, session)
    return l_out


@router.post("/from-conversation/{conversation_id}", response_model=LeadOut, status_code=status.HTTP_201_CREATED)
async def create_lead_from_conversation(
    conversation_id: str,
    _: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadOut:
    conversation = await session.get(Conversation, conversation_id)
    if not conversation:
        raise HTTPException(status_code=404, detail="Conversation not found")

    contact = await session.get(Contact, conversation.contact_id)
    inbox = await session.get(Inbox, conversation.inbox_id)
    if not contact or not inbox:
        raise HTTPException(status_code=404, detail="Conversation contact or inbox not found")

    lead = await upsert_lead_from_inbound_message(
        session,
        inbox=inbox,
        contact=contact,
        conversation=conversation,
        channel_type=conversation.channel_type,
        capture_source="manual",
    )
    await session.commit()
    await session.refresh(lead)
    return LeadOut.model_validate(lead)


@router.get("/check-phone")
async def check_phone_unique(
    phone: str = Query(..., min_length=1),
    exclude_lead_id: str | None = Query(default=None),
    _: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
):
    phone_stripped = phone.strip()
    if not phone_stripped:
        return {"available": True, "existing_name": None}
    
    filters = [Lead.phone == phone_stripped]
    if exclude_lead_id:
        filters.append(Lead.id != exclude_lead_id)
        
    q = select(Lead).where(*filters)
    res = await session.execute(q)
    existing = res.scalars().first()
    if existing:
        existing_name = existing.company_name or existing.contact_person or "Unnamed Lead"
        return {"available": False, "existing_name": existing_name}
    return {"available": True, "existing_name": None}


@router.get("/{lead_id}", response_model=LeadOut)
async def get_lead(
    lead_id: str,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadOut:
    lead = await _get_lead_or_404(lead_id, session, auth)
    l_out = LeadOut.model_validate(lead)
    l_out.assigned_user_ids = await _populate_lead_assignments(lead, session)
    return l_out


@router.patch("/{lead_id}", response_model=LeadOut)
async def update_lead(
    lead_id: str,
    payload: LeadUpdate,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadOut:
    lead = await _get_lead_or_404(lead_id, session, auth)
    changes = payload.model_dump(exclude_unset=True)
    if changes.get("priority"):
        changes["priority"] = changes["priority"].lower()
    if "tags" in changes:
        changes["tags"] = _normalize_tags(changes.get("tags"))
    await _validate_config_ids(changes, session)

    if "phone" in changes:
        phone = changes.get("phone")
        if phone and phone.strip():
            phone_stripped = phone.strip()
            q = select(Lead).where(Lead.phone == phone_stripped, Lead.id != lead_id)
            res = await session.execute(q)
            existing = res.scalars().first()
            if existing:
                existing_name = existing.company_name or existing.contact_person or "Unnamed Lead"
                raise HTTPException(
                    status_code=400,
                    detail=f"Lead with this phone number already exists: {existing_name}",
                )

    previous_status = lead.status
    assigned_user_ids = changes.pop("assigned_user_ids", None)

    assignment_changed = False
    if "assigned_user_id" in changes and changes["assigned_user_id"] != lead.assigned_user_id:
        assignment_changed = True
    if "assigned_agent_id" in changes and changes["assigned_agent_id"] != lead.assigned_agent_id:
        assignment_changed = True
    if assigned_user_ids is not None:
        current_assigned_uids = set(await _populate_lead_assignments(lead, session))
        if set(assigned_user_ids) != current_assigned_uids:
            assignment_changed = True

    for key, value in changes.items():
        setattr(lead, key, value)

    is_assigned = (lead.assigned_user_id or lead.assigned_agent_id or (assigned_user_ids if assigned_user_ids is not None else await _populate_lead_assignments(lead, session)))

    if assignment_changed and is_assigned:
        from datetime import datetime, timezone
        lead.assigned_at = datetime.now(timezone.utc)

    if not is_assigned:
        lead.assigned_at = None
        lead.target_stage_by = None
    elif assignment_changed or "sla_duration_hours" in changes:
        if lead.assigned_at and lead.sla_duration_hours:
            from datetime import timedelta
            lead.target_stage_by = lead.assigned_at + timedelta(hours=lead.sla_duration_hours)
        else:
            lead.target_stage_by = None

    if changes and "untouched" not in changes:
        lead.untouched = False

    if "status" in changes:
        await record_stage_history(
            session,
            workspace_id=auth.user_id,
            lead=lead,
            old_stage=previous_status,
            new_stage=lead.status,
            changed_by_user_id=auth.user_id,
            change_reason="manual_update",
        )

    if assigned_user_ids is not None:
        from app.db.models import Assignment
        from sqlalchemy import delete
        await session.execute(delete(Assignment).where(Assignment.lead_id == lead.id))
        for uid in assigned_user_ids:
            session.add(Assignment(
                organization_id=auth.org_id,
                lead_id=lead.id,
                user_id=uid,
                assigned_by_id=auth.user_id
            ))

    text_blob = " ".join(filter(None, [lead.notes, lead.raw_note, lead.address, lead.industry]))
    await record_preference_history(
        session,
        workspace_id=auth.user_id,
        lead=lead,
        text=text_blob,
        detected_from="lead_update",
    )
    await evaluate_intelligence_alerts(
        session,
        workspace_id=auth.user_id,
        lead=lead,
        text=text_blob,
        source="lead_update",
    )

    # Sync changes to Customer if converted
    if lead.converted_customer_id:
        customer_obj = await session.get(Customer, lead.converted_customer_id)
        if customer_obj:
            if "status" in changes or "lead_stage_id" in changes:
                stage_name = lead.status
                if lead.lead_stage_id:
                    from app.db.models import LeadStage
                    stage_ent = await session.get(LeadStage, lead.lead_stage_id)
                    if stage_ent:
                        stage_name = stage_ent.name
                customer_obj.stage = stage_name
            
            if "company_name" in changes:
                customer_obj.company_name = lead.company_name
            if "contact_person" in changes:
                customer_obj.contact_person = lead.contact_person
            if "phone" in changes:
                customer_obj.phone = lead.phone
            if "address" in changes:
                customer_obj.address = lead.address
            if "last_education" in changes:
                customer_obj.last_education = lead.last_education
            if lead.industry_data and "countries_applied" in lead.industry_data:
                customer_obj.countries_applied = lead.industry_data["countries_applied"]

    await session.commit()
    await session.refresh(lead)
    l_out = LeadOut.model_validate(lead)
    l_out.assigned_user_ids = await _populate_lead_assignments(lead, session)
    return l_out


@router.delete("/{lead_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_lead(
    lead_id: str,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> None:
    lead = await _get_lead_or_404(lead_id, session, auth)
    if auth.role != "super_admin" and not await has_permission(session, auth.user_id, auth, "leads.manage"):
        raise HTTPException(status_code=403, detail="Permission denied")
    await session.delete(lead)
    await session.commit()


@router.get("/{lead_id}/activities", response_model=list[LeadActivityOut])
async def list_lead_activities(
    lead_id: str,
    activity_type: str | None = Query(default=None),
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> list[LeadActivityOut]:
    lead = await _get_lead_or_404(lead_id, session, auth)
    query = select(LeadActivity).where(LeadActivity.lead_id == lead_id)
    if activity_type:
        query = query.where(LeadActivity.activity_type == activity_type)
    query = query.order_by(LeadActivity.created_at.desc())
    rows = (await session.execute(query)).scalars().all()
    return [LeadActivityOut.model_validate(row) for row in rows]


@router.get("/{lead_id}/timeline", response_model=list[LeadTimelineItem])
async def get_lead_timeline(
    lead_id: str,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> list[LeadTimelineItem]:
    lead = await _get_lead_or_404(lead_id, session, auth)
    items: list[LeadTimelineItem] = [
        LeadTimelineItem(
            id=activity.id,
            item_type="activity",
            activity_type=activity.activity_type,
            direction=activity.direction,
            platform=activity.platform,
            title=activity.title,
            content=activity.content,
            created_by_user_id=activity.created_by_user_id,
            due_at=activity.due_at,
            completed_at=activity.completed_at,
            created_at=activity.created_at,
        )
        for activity in (
            await session.execute(select(LeadActivity).where(LeadActivity.lead_id == lead_id))
        ).scalars().all()
    ]

    if lead.conversation_id:
        messages = (
            await session.execute(
                select(Message)
                .where(Message.conversation_id == lead.conversation_id)
                .order_by(Message.created_at.desc())
            )
        ).scalars().all()
        items.extend(
            LeadTimelineItem(
                id=message.id,
                item_type="message",
                activity_type="message",
                direction=message.message_type,
                platform=lead.capture_source or "conversation",
                title="Conversation message",
                content=message.content,
                created_by_user_id=message.sender_id,
                created_at=message.created_at,
            )
            for message in messages
        )

    return sorted(items, key=lambda item: item.created_at, reverse=True)


@router.post("/{lead_id}/activities", response_model=LeadActivityOut, status_code=status.HTTP_201_CREATED)
async def create_lead_activity(
    lead_id: str,
    payload: LeadActivityCreate,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadActivityOut:
    lead = await _get_lead_or_404(lead_id, session, auth)
    data = payload.model_dump(by_alias=False)
    metadata = data.pop("metadata", None)
    activity = LeadActivity(
        **data,
        lead_id=lead_id,
        created_by_user_id=auth.user_id,
        activity_metadata=metadata or {},
    )
    lead.untouched = False
    if payload.activity_type == "follow_up" and payload.due_at:
        lead.follow_up_date = payload.due_at
    session.add(activity)
    await session.commit()
    await session.refresh(activity)
    return LeadActivityOut.model_validate(activity)


@router.post("/{lead_id}/share-links", response_model=LeadShareOut, status_code=status.HTTP_201_CREATED)
async def create_lead_share_link(
    lead_id: str,
    payload: LeadShareCreate,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadShareOut:
    await _get_lead_or_404(lead_id, session, auth)
    if auth.role != "super_admin" and not await _can_view_all_leads(auth, session):
        raise HTTPException(status_code=403, detail="Permission denied")
    org = None
    if auth.org_id:
        org = await session.get(Organization, auth.org_id)

    if payload.mode == "public" and org and not org.allow_public_shares:
        raise HTTPException(status_code=403, detail="Public share links are disabled")

    allowed_emails = _normalize_emails(payload.allowed_emails)
    if payload.mode == "restricted" and not allowed_emails:
        raise HTTPException(status_code=400, detail="Allowed emails are required for restricted sharing")

    expires_at = None
    if payload.expires_in_days:
        expires_at = datetime.utcnow() + timedelta(days=payload.expires_in_days)
    elif org and org.default_share_expiry_days:
        expires_at = datetime.utcnow() + timedelta(days=org.default_share_expiry_days)

    token = secrets.token_urlsafe(24)
    share = LeadShareLink(
        lead_id=lead_id,
        created_by_user_id=auth.user_id,
        token=token,
        is_public=(payload.mode == "public"),
        allowed_emails=allowed_emails,
        expires_at=expires_at,
    )
    session.add(share)
    await session.commit()
    await session.refresh(share)

    share_url = f"{settings.frontend_origin.rstrip('/')}/share/lead/{share.token}"
    return LeadShareOut(
        id=share.id,
        lead_id=share.lead_id,
        token=share.token,
        share_url=share_url,
        is_public=share.is_public,
        allowed_emails=share.allowed_emails or [],
        expires_at=share.expires_at,
        created_at=share.created_at,
    )


@router.patch("/{lead_id}/activities/{activity_id}", response_model=LeadActivityOut)
async def update_lead_activity(
    lead_id: str,
    activity_id: str,
    payload: LeadActivityUpdate,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> LeadActivityOut:
    lead = await _get_lead_or_404(lead_id, session, auth)
    activity = await session.get(LeadActivity, activity_id)
    if not activity or activity.lead_id != lead_id:
        raise HTTPException(status_code=404, detail="Lead activity not found")
    changes = payload.model_dump(exclude_unset=True, by_alias=False)
    metadata = changes.pop("metadata", None)
    for key, value in changes.items():
        setattr(activity, key, value)
    if metadata is not None:
        activity.activity_metadata = metadata
    await session.commit()
    await session.refresh(activity)
    return LeadActivityOut.model_validate(activity)


@router.delete("/{lead_id}/activities/{activity_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_lead_activity(
    lead_id: str,
    activity_id: str,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
) -> None:
    lead = await _get_lead_or_404(lead_id, session, auth)
    activity = await session.get(LeadActivity, activity_id)
    if not activity or activity.lead_id != lead_id:
        raise HTTPException(status_code=404, detail="Lead activity not found")
    await session.delete(activity)
    await session.commit()


@router.post("/{lead_id}/convert", status_code=status.HTTP_201_CREATED)
async def convert_lead(
    lead_id: str,
    payload: LeadConvertPayload,
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
):
    lead = await _get_lead_or_404(lead_id, session, auth)
    if auth.role != "super_admin" and not await _can_view_all_leads(auth, session):
        raise HTTPException(status_code=403, detail="Permission denied")
    if lead.converted_customer_id:
        customer = await session.get(Customer, lead.converted_customer_id)
        if customer:
            return {"customer": CustomerOut.model_validate(customer).model_dump(), "invoice": None, "opportunity_id": None}

    # Get stage name
    stage_name = lead.status
    if lead.lead_stage_id:
        from app.db.models import LeadStage
        stage_ent = await session.get(LeadStage, lead.lead_stage_id)
        if stage_ent:
            stage_name = stage_ent.name

    customer = Customer(
        organization_id=lead.organization_id,
        branch_id=lead.branch_id,
        company_name=lead.company_name,
        contact_person=lead.contact_person,
        email=lead.email,
        phone=lead.phone,
        address=lead.address,
        assigned_user_id=lead.assigned_user_id or auth.user_id,
        stage=stage_name,
        group_name=lead.source,
        tags={"lead_id": lead.id, "lead_source": lead.source},
        last_education=lead.last_education,
        countries_applied=lead.industry_data.get("countries_applied", []) if lead.industry_data else [],
        notes=f"Converted from lead {lead.id}. Source: {lead.source or 'unsourced' }.",
    )
    session.add(customer)
    await session.flush()

    budget = payload.budget or 0

    # Determine first stage of active pipeline for organization
    from app.db.models import Pipeline, PipelineStage
    pipe_res = await session.execute(
        select(Pipeline).where(Pipeline.organization_id == lead.organization_id, Pipeline.is_active == True)
    )
    pipeline = pipe_res.scalars().first()
    first_stage_name = "discovery" # default fallback
    if pipeline:
        stage_res = await session.execute(
            select(PipelineStage)
            .where(PipelineStage.pipeline_id == pipeline.id)
            .order_by(PipelineStage.position.asc())
        )
        first_stage = stage_res.scalars().first()
        if first_stage:
            first_stage_name = first_stage.name

    opportunity = Opportunity(
        organization_id=lead.organization_id,
        branch_id=lead.branch_id,
        customer_id=customer.id,
        title=f"Opportunity for {customer.company_name or customer.contact_person or 'Customer'}",
        stage=first_stage_name,
        estimated_value=budget,
        currency="BDT",
    )
    session.add(opportunity)
    await session.flush()

    sales_order = SalesOrder(
        organization_id=lead.organization_id,
        branch_id=lead.branch_id,
        customer_id=customer.id,
        handler_user_id=customer.assigned_user_id,
        status="draft",
        payment_status="pending",
        total_amount=budget,
        currency="BDT",
        remark=f"Initial invoice for {customer.company_name}",
    )
    session.add(sales_order)
    await session.flush()

    task = Task(
        organization_id=lead.organization_id,
        branch_id=lead.branch_id,
        entity_type="opportunity",
        entity_id=opportunity.id,
        assigned_user_id=customer.assigned_user_id,
        title="Initial Outreach for New Deal",
        description="Follow up on the newly converted lead to kick off the discovery process.",
        priority="high",
    )
    session.add(task)

    audit = AuditLog(
        entity_type="lead",
        entity_id=lead.id,
        action="converted_to_customer",
        previous_value={"status": lead.status},
        new_value={"status": "won", "customer_id": customer.id, "opportunity_id": opportunity.id},
        performed_by_user_id=auth.user_id,
    )
    session.add(audit)

    old_status = lead.status
    lead.converted_customer_id = customer.id
    lead.status = "won"
    lead.untouched = False
    await record_stage_history(
        session,
        workspace_id=auth.user_id,
        lead=lead,
        old_stage=old_status,
        new_stage="won",
        changed_by_user_id=auth.user_id,
        change_reason="conversion",
    )

    await session.commit()
    await session.refresh(customer)
    await session.refresh(sales_order)
    return {
        "customer": CustomerOut.model_validate(customer).model_dump(),
        "invoice": {
            "id": sales_order.id,
            "customer_id": sales_order.customer_id,
            "status": sales_order.status,
            "payment_status": sales_order.payment_status,
            "total_amount": float(sales_order.total_amount),
            "currency": sales_order.currency,
            "remark": sales_order.remark,
            "created_at": sales_order.created_at.isoformat(),
        },
        "opportunity_id": opportunity.id,
    }


@router.get("/export")
async def export_leads(
    status_filter: str | None = Query(default=None, alias="status"),
    stage_id: str | None = Query(default=None),
    source_id: str | None = Query(default=None),
    tag: str | None = Query(default=None),
    priority: str | None = Query(default=None),
    search: str | None = Query(default=None),
    start_date: datetime | None = Query(default=None),
    end_date: datetime | None = Query(default=None),
    assigned_user_id: str | None = Query(default=None),
    branch_id: str | None = Query(default=None),
    quick_filter: str | None = Query(default=None),
    lead_ids: list[str] | None = Query(default=None),
    format: str = Query(default="excel", pattern="^(csv|excel)$"),
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
):
    if not await _can_view_leads(auth, session):
        raise HTTPException(status_code=403, detail="Permission denied")

    filters = []

    if lead_ids:
        filters.append(Lead.id.in_(lead_ids))

    if status_filter and status_filter != "all":
        filters.append(Lead.status == status_filter)
    if branch_id:
        filters.append(Lead.branch_id == branch_id)
    if stage_id:
        filters.append(Lead.lead_stage_id == stage_id)
    if source_id:
        filters.append(Lead.lead_source_id == source_id)
    if tag:
        tag_value = tag.strip().lower()
        if tag_value:
            filters.append(Lead.tags.contains([tag_value]))
    if priority and priority != "all":
        filters.append(Lead.priority == priority.lower())
    if search:
        needle = f"%{search.strip()}%"
        filters.append(
            or_(
                Lead.company_name.ilike(needle),
                Lead.contact_person.ilike(needle),
                Lead.phone.ilike(needle),
                Lead.email.ilike(needle),
            )
        )
    if start_date:
        filters.append(Lead.created_at >= start_date)
    if end_date:
        filters.append(Lead.created_at <= end_date)
    if assigned_user_id:
        filters.append(or_(Lead.assigned_user_id == assigned_user_id, Lead.assigned_agent_id == assigned_user_id))
    if quick_filter == "assigned_to_me":
        filters.append(or_(Lead.assigned_user_id == auth.user_id, Lead.assigned_agent_id == auth.user_id))
    elif quick_filter == "untouched":
        filters.append(Lead.untouched.is_(True))
    elif quick_filter == "followups_due":
        filters.append(and_(Lead.follow_up_date.is_not(None), Lead.follow_up_date <= func.now()))

    query = select(Lead)
    if filters:
        query = query.where(*filters)

    from app.api.deps import apply_tenant_filters
    query = apply_tenant_filters(query, auth, Lead)
    query = query.order_by(Lead.created_at.desc())

    rows = (await session.execute(query)).scalars().all()

    if format == "excel":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"

        headers = [
            "Company Name", "Contact Person", "Phone", "Email", 
            "Address", "Last Education", "Priority", "Status", "Source", "Notes"
        ]
        ws.append(headers)

        for lead in rows:
            ws.append([
                lead.company_name or "",
                lead.contact_person or "",
                lead.phone or "",
                lead.email or "",
                lead.address or "",
                lead.last_education or "",
                lead.priority or "medium",
                lead.status or "new",
                lead.source or "",
                lead.notes or ""
            ])

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)

        return StreamingResponse(
            out_buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=leads_export.xlsx"}
        )

    # Fallback to CSV
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Company Name", "Contact Person", "Phone", "Email", 
        "Address", "Last Education", "Priority", "Status", "Source", "Notes"
    ])

    for lead in rows:
        writer.writerow([
            lead.company_name or "",
            lead.contact_person or "",
            lead.phone or "",
            lead.email or "",
            lead.address or "",
            lead.last_education or "",
            lead.priority or "medium",
            lead.status or "new",
            lead.source or "",
            lead.notes or ""
        ])

    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=leads_export.csv"}
    )


@router.post("/bulk-upload")
async def bulk_upload_leads(
    file: UploadFile = File(...),
    auth: AuthContext = Depends(get_current_auth),
    session: AsyncSession = Depends(get_session_dep),
):
    if auth.role != "super_admin" and not await has_permission(session, auth.user_id, auth, "leads.manage"):
        raise HTTPException(status_code=403, detail="Permission denied")

    content = await file.read()
    is_excel = file.filename.endswith(".xlsx") or file.filename.endswith(".xls") or file.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    if is_excel:
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = wb.active
            rows_data = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(exc)}")

        if not rows_data or not rows_data[0]:
            raise HTTPException(status_code=400, detail="Empty Excel file.")

        raw_headers = [str(h).strip() if h is not None else "" for h in rows_data[0]]
        rows_list = []
        for r in rows_data[1:]:
            row_dict = {}
            for idx, val in enumerate(r):
                if idx < len(raw_headers):
                    row_dict[raw_headers[idx]] = str(val).strip() if val is not None else ""
            rows_list.append(row_dict)

        class ExcelReader:
            def __init__(self, fieldnames, rows):
                self.fieldnames = fieldnames
                self.rows = rows
            def __iter__(self):
                return iter(self.rows)

        reader = ExcelReader(raw_headers, rows_list)
    else:
        try:
            csv_text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                csv_text = content.decode("latin1")
            except Exception:
                raise HTTPException(status_code=400, detail="Unable to decode file. Please upload a valid UTF-8, Latin1 CSV, or Excel file.")

        reader = csv.DictReader(io.StringIO(csv_text))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="Empty CSV headers.")

    headers = [h.strip().lower() for h in reader.fieldnames]

    name_field = None
    for option in ["company name", "company/client name", "company_name", "name", "client name", "client_name"]:
        if option in headers:
            name_field = reader.fieldnames[headers.index(option)]
            break

    if not name_field:
        raise HTTPException(status_code=400, detail="CSV must contain a 'Company Name' or 'Name' column.")

    phone_field = None
    email_field = None
    address_field = None
    education_field = None
    priority_field = None
    status_field = None
    source_field = None
    notes_field = None
    contact_person_field = None

    for idx, h in enumerate(headers):
        field_lower = h.strip()
        if field_lower in ["phone", "phone number", "phone_number"]:
            phone_field = reader.fieldnames[idx]
        elif field_lower in ["email", "email address", "email_address"]:
            email_field = reader.fieldnames[idx]
        elif field_lower in ["address", "location"]:
            address_field = reader.fieldnames[idx]
        elif field_lower in ["last education", "education", "last_education"]:
            education_field = reader.fieldnames[idx]
        elif field_lower in ["priority"]:
            priority_field = reader.fieldnames[idx]
        elif field_lower in ["status", "stage", "lead stage", "lead_stage"]:
            status_field = reader.fieldnames[idx]
        elif field_lower in ["source", "lead source", "lead_source"]:
            source_field = reader.fieldnames[idx]
        elif field_lower in ["notes", "note", "description"]:
            notes_field = reader.fieldnames[idx]
        elif field_lower in ["contact person", "contact_person"]:
            contact_person_field = reader.fieldnames[idx]

    success_count = 0
    errors = []

    # Get existing phone numbers in organization
    existing_phones_res = await session.execute(
        select(Lead.phone).where(Lead.organization_id == auth.org_id, Lead.phone.is_not(None))
    )
    existing_phones = set(existing_phones_res.scalars().all())

    # Get pipeline stages for the organization
    from app.db.models import Pipeline, PipelineStage
    org_stages_res = await session.execute(
        select(PipelineStage)
        .join(Pipeline)
        .where(Pipeline.organization_id == auth.org_id)
    )
    org_stages = list(org_stages_res.scalars().all())
    org_stage_map = {s.name.lower().strip(): s for s in org_stages}

    row_idx = 1
    for row in reader:
        row_idx += 1
        company_name = row.get(name_field)
        if not company_name or not company_name.strip():
            errors.append(f"Row {row_idx}: Missing company/client name")
            continue

        phone = row.get(phone_field).strip() if phone_field and row.get(phone_field) else None
        if phone and phone in existing_phones:
            errors.append(f"Row {row_idx}: Phone number '{phone}' already exists")
            continue

        email = row.get(email_field).strip() if email_field and row.get(email_field) else None
        address = row.get(address_field).strip() if address_field and row.get(address_field) else None
        education = row.get(education_field).strip() if education_field and row.get(education_field) else None
        priority = row.get(priority_field).strip().lower() if priority_field and row.get(priority_field) else "medium"
        if priority not in ["low", "medium", "high"]:
            priority = "medium"

        status_val = row.get(status_field).strip() if status_field and row.get(status_field) else "new"
        source = row.get(source_field).strip() if source_field and row.get(source_field) else None
        notes = row.get(notes_field).strip() if notes_field and row.get(notes_field) else None
        contact_person = row.get(contact_person_field).strip() if contact_person_field and row.get(contact_person_field) else None

        lead_stage_id = None
        status_val_lower = status_val.lower().strip()
        if status_val_lower in org_stage_map:
            lead_stage_id = org_stage_map[status_val_lower].id
            status_val = org_stage_map[status_val_lower].name

        lead_obj = Lead(
            organization_id=auth.org_id,
            branch_id=auth.branch_id,
            company_name=company_name.strip(),
            contact_person=contact_person,
            phone=phone,
            email=email,
            address=address,
            last_education=education,
            priority=priority,
            status=status_val,
            lead_stage_id=lead_stage_id,
            source=source,
            notes=notes,
            untouched=True
        )
        session.add(lead_obj)
        if phone:
            existing_phones.add(phone)
        success_count += 1

    await session.commit()
    return {"success": True, "imported_count": success_count, "errors": errors}
